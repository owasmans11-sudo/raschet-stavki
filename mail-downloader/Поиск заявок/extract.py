"""Достаём текст из файлов: PDF, Word, Excel, картинки (OCR), архивы и прочее."""

import io
import os
import re
import shutil
import tempfile
import zipfile

# Tesseract сам запускает несколько потоков; мы и так распознаём файлы параллельно
os.environ.setdefault("OMP_THREAD_LIMIT", "1")

MAX_CHARS = 30000          # больше текста для поиска заявки не нужно
OCR_PDF_PAGES = 3          # сканы PDF: распознаём только первые страницы
MIN_IMAGE_BYTES = 25_000   # мелкие картинки (логотипы, подписи) пропускаем
MIN_IMAGE_SIDE = 500

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".heic"}
ARCHIVE_EXT = {".zip", ".rar", ".7z"}
TEXT_EXT = {".txt", ".csv", ".xml", ".htm", ".html", ".eml", ".json"}


# ---------------- OCR ----------------
_tess = {"checked": False, "cmd": None, "langs": "rus+eng", "tessdata": None}


def find_tesseract():
    if _tess["checked"]:
        return _tess["cmd"]
    _tess["checked"] = True
    cands = [shutil.which("tesseract"),
             r"C:\Program Files\Tesseract-OCR\tesseract.exe",
             r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
             os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe")]
    for c in cands:
        if c and os.path.isfile(c):
            _tess["cmd"] = c
            break
    # Свой tessdata с русским языком, если в системной установке его нет
    own = os.path.join(os.path.expanduser("~"), ".zayavki", "tessdata")
    if os.path.isfile(os.path.join(own, "rus.traineddata")):
        _tess["tessdata"] = own
    return _tess["cmd"]


def ocr_image(img):
    """img: PIL.Image. Возвращает распознанный текст или ''."""
    cmd = find_tesseract()
    if not cmd:
        return ""
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = cmd
    cfg = "--psm 3"
    if _tess["tessdata"]:
        cfg += f' --tessdata-dir "{_tess["tessdata"]}"'
    img = img.convert("L")
    w, h = img.size
    if max(w, h) > 3500:                       # слишком большие уменьшаем, так быстрее
        k = 3500 / max(w, h)
        img = img.resize((int(w * k), int(h * k)))
    try:
        return pytesseract.image_to_string(img, lang=_tess["langs"], config=cfg,
                                           timeout=120)
    except Exception:
        return ""


def from_image_bytes(data):
    from PIL import Image, ImageOps
    if len(data) < MIN_IMAGE_BYTES:
        return "", "мелкая картинка"
    img = Image.open(io.BytesIO(data))
    try:
        img.seek(0)
    except Exception:
        pass
    img = ImageOps.exif_transpose(img)
    if max(img.size) < MIN_IMAGE_SIDE:
        return "", "мелкая картинка"
    if not find_tesseract():
        return "", "нет OCR"
    return ocr_image(img), "OCR"


# ---------------- форматы ----------------
def from_pdf(data):
    import pymupdf as fitz
    doc = fitz.open(stream=data, filetype="pdf")
    parts = []
    for page in doc:
        parts.append(page.get_text())
        if sum(map(len, parts)) > MAX_CHARS:
            break
    text = "\n".join(parts)
    if len(text.strip()) > 80:
        return text, "PDF"
    # Скан без текстового слоя: распознаём первые страницы
    if not find_tesseract():
        return text, "PDF-скан, нет OCR"
    from PIL import Image
    parts = []
    for i, page in enumerate(doc):
        if i >= OCR_PDF_PAGES:
            break
        pix = page.get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        parts.append(ocr_image(img))
    return "\n".join(parts), "PDF-скан OCR"


def from_docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                if not cells or cells[-1] != c.text:   # объединённые ячейки повторяются
                    cells.append(c.text)
            parts.append(" | ".join(cells))
    for s in d.sections:
        for p in s.header.paragraphs:
            parts.append(p.text)
    text = "\n".join(parts)
    # Заявка часто вставлена в Word картинкой
    if len(text.strip()) < 80 and find_tesseract():
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for n in z.namelist():
                if n.startswith("word/media/") and os.path.splitext(n)[1].lower() in IMAGE_EXT:
                    t, _ = from_image_bytes(z.read(n))
                    text += "\n" + t
        return text, "Word + OCR"
    return text, "Word"


def from_xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, n = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v not in (None, "")]
            if vals:
                line = " | ".join(vals)
                parts.append(line)
                n += len(line)
            if n > MAX_CHARS:
                break
    return "\n".join(parts), "Excel"


def from_xls(data):
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    parts, n = [], 0
    for sh in wb.sheets():
        for r in range(sh.nrows):
            vals = [str(v) for v in sh.row_values(r) if v not in (None, "")]
            if vals:
                line = " | ".join(vals)
                parts.append(line)
                n += len(line)
            if n > MAX_CHARS:
                break
    return "\n".join(parts), "Excel"


def from_doc(data):
    """Старый .doc: текст лежит в UTF-16LE или cp1251, достаём читаемые куски."""
    u = re.findall(r"(?:[\u0020-\u007e\u0400-\u04ff\u2116\u00ab\u00bb\r\n\t]){4,}",
                   data.decode("utf-16le", errors="ignore"))
    a = re.findall(rb"[\x20-\x7e\xc0-\xff\xa8\xb8\xb9\r\n\t]{6,}", data)
    text = "\n".join(u) + "\n" + "\n".join(x.decode("cp1251", "ignore") for x in a)
    return text, "Word (doc)"


def from_rtf(data):
    from striprtf.striprtf import rtf_to_text
    raw = data.decode("cp1251", errors="ignore")
    return rtf_to_text(raw, errors="ignore"), "RTF"


def from_pptx(data):
    from pptx import Presentation
    p = Presentation(io.BytesIO(data))
    parts = []
    for s in p.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
    return "\n".join(parts), "PowerPoint"


def from_odt(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml").decode("utf-8", "ignore")
    return re.sub(r"<[^>]+>", " ", xml), "OpenDocument"


def from_text(data):
    for enc in ("utf-8", "cp1251"):
        try:
            t = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        t = data.decode("utf-8", "ignore")
    return re.sub(r"<[^>]+>", " ", t), "Текст"


HANDLERS = {
    ".pdf": from_pdf, ".docx": from_docx, ".docm": from_docx, ".dotx": from_docx,
    ".xlsx": from_xlsx, ".xlsm": from_xlsx, ".xls": from_xls,
    ".doc": from_doc, ".rtf": from_rtf, ".pptx": from_pptx,
    ".odt": from_odt, ".ods": from_odt,
}


def extract_bytes(name, data):
    """Текст и способ извлечения для одного файла (не архива)."""
    ext = os.path.splitext(name)[1].lower()
    if ext in IMAGE_EXT:
        text, how = from_image_bytes(data)
    elif ext in HANDLERS:
        text, how = HANDLERS[ext](data)
    elif ext in TEXT_EXT:
        text, how = from_text(data)
    elif data[:4] == b"%PDF":
        text, how = from_pdf(data)
    else:
        return "", "формат не поддерживается"
    return (text or "")[:MAX_CHARS], how


# ---------------- архивы ----------------
def archive_members(path):
    """Отдаёт (имя_внутри, байты) для каждого файла архива. Вложенные архивы тоже."""
    ext = os.path.splitext(path)[1].lower()
    with tempfile.TemporaryDirectory() as tmp:
        if ext == ".zip":
            with zipfile.ZipFile(path) as z:
                for info in z.infolist():
                    if info.is_dir() or info.file_size > 200_000_000:
                        continue
                    yield _fix_zip_name(info), z.read(info)
            return
        if ext == ".7z":
            import py7zr
            with py7zr.SevenZipFile(path) as z:
                z.extractall(tmp)
        elif ext == ".rar":
            import rarfile
            _setup_unrar(rarfile)
            with rarfile.RarFile(path) as z:
                z.extractall(tmp)
        for root, _, files in os.walk(tmp):
            for f in files:
                full = os.path.join(root, f)
                with open(full, "rb") as fh:
                    yield os.path.relpath(full, tmp), fh.read()


def _fix_zip_name(info):
    # Архивы из Windows часто пишут русские имена в cp866 без флага UTF-8
    name = info.filename
    if not info.flag_bits & 0x800:
        try:
            name = name.encode("cp437").decode("cp866")
        except Exception:
            pass
    return name


def _setup_unrar(rarfile):
    for c in (r"C:\Program Files\7-Zip\7z.exe", r"C:\Program Files (x86)\7-Zip\7z.exe",
              shutil.which("7z") or "", r"C:\Program Files\WinRAR\UnRAR.exe",
              shutil.which("unrar") or ""):
        if c and os.path.isfile(c):
            if c.lower().endswith("7z.exe") or c.endswith("7z"):
                rarfile.SEVENZIP_TOOL = c
            else:
                rarfile.UNRAR_TOOL = c
            return
